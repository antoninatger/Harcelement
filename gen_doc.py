# -*- coding: utf-8 -*-
"""
gen_doc.py - produit le document partenaires à partir du jeu lui-même.

    node dump_textes.js      (écrit _textes_dump.json depuis les fichiers de jeu)
    python gen_doc.py        (écrit doc_textes_partenaires.docx et .html)

Remplace l'ancien gen_docx.py, où les 9 500 mots du jeu étaient recopiés à la
main : le document dérivait dès la première modification des textes. Ici, la
mise en forme reste écrite ici, le contenu vient de _modele.py, qui lit le jeu.
"""
import html as _h
import importlib.util, os, sys

ICI = os.path.dirname(os.path.abspath(__file__)) or '.'
os.chdir(ICI)

def _charger(nom, chemin):
    spec = importlib.util.spec_from_file_location(nom, chemin)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[nom] = mod
    spec.loader.exec_module(mod)
    return mod

if not os.path.exists('_textes_dump.json'):
    sys.exit("_textes_dump.json manquant - lancez d'abord : node dump_textes.js")

socle = _charger('_socle_docx', '_socle_docx.py')   # styles Word + helpers
modele = _charger('_modele', '_modele.py')          # le plan, tiré du jeu
B = modele.B

# ═══════════════════════════════════════════════════════════════════════════
#  RENDU WORD
# ═══════════════════════════════════════════════════════════════════════════
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
s = socle
doc = s.doc

COUL = {'good': s.C_GREEN, 'ok': s.C_ORANGE, 'bad': s.C_RED, 'neutral': s.C_GREY}

def w_cover(titre, sous, meta):
    ey = doc.add_paragraph(); ey.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.set_spacing(ey, before=560, after=100)
    s.add_run(ey, 'Jeu pédagogique sur le harcèlement scolaire', size=8, color=s.C_DIM, font='Arial')
    ti = doc.add_paragraph(); ti.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.set_spacing(ti, before=60, after=60); s.bottom_border(ti, '1a1a1a', sz=12)
    s.add_run(ti, titre, italic=True, size=26)
    su = doc.add_paragraph(); su.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.set_spacing(su, before=80, after=40)
    s.add_run(su, sous, size=10.5, color=s.C_GREY)
    me = doc.add_paragraph(); me.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.set_spacing(me, before=40, after=200)
    s.add_run(me, meta, size=8, color=s.C_DIM, font='Arial')

def w_table(entetes, lignes):
    t = doc.add_table(rows=1, cols=len(entetes)); t.style = 'Table Grid'
    for i, e in enumerate(entetes):
        c = t.rows[0].cells[i]; c.text = ''
        s.add_run(c.paragraphs[0], e, bold=True, size=9, font='Arial')
    for ln in lignes:
        cells = t.add_row().cells
        for i, v in enumerate(ln):
            cells[i].text = ''
            s.add_run(cells[i].paragraphs[0], v, size=9.5)
    doc.add_paragraph()

def rendu_docx():
    for b in B:
        k = b[0]
        if k == 'cover':    w_cover(b[1], b[2], b[3])
        elif k == 'h1':     s.part_header(b[1], b[2])
        elif k == 'h2':     s.heading2(b[1])
        elif k == 'h3':     s.heading3(b[1])
        elif k == 'p':      s.p(b[1], size=10.5, after=70)
        elif k == 'note':   s.p(b[1], size=9, after=70, italic=True, color=s.C_GREY)
        elif k == 'bullet': s.bullet(b[1])
        elif k == 'table':  w_table(b[1], b[2])
        elif k == 'sep':    s.time_sep(b[1])
        elif k == 'prompt':
            para = doc.add_paragraph(); s.set_spacing(para, before=60, after=20)
            s.add_run(para, b[1], italic=True, size=9, color=s.C_DIM, font='Arial')
        elif k == 'pensee':
            para = doc.add_paragraph(); s.set_spacing(para, before=40, after=40)
            s.set_indent(para, left=200); s.shading(para, 'F2F8F2'); s.left_border(para, '4dcc70', sz=10)
            s.add_run(para, b[1], italic=True, size=9.5, color=s.C_GREEN)
        elif k == 'dial':
            cle = b[1] if b[1] in s.SPK else 'narr'
            s.dialogue_line(cle, b[2], b[3], italic_text=(len(b) > 4 and b[4]))
        elif k == 'choix':
            kind, lettre, txt, note, rep, qui = b[1], b[2], b[3], b[4], b[5], b[6]
            para = doc.add_paragraph(); s.set_spacing(para, before=20, after=20)
            s.set_indent(para, left=700); s.shading(para, s.CHOICE_BG[kind]); s.left_border(para, s.CHOICE_HEX[kind], sz=14)
            s.add_run(para, '%s  ' % lettre, bold=True, size=9, color=COUL[kind], font='Arial')
            s.add_run(para, txt, size=10)
            if note:
                para.add_run('\n'); s.add_run(para, note, italic=True, size=8.5, color=s.C_DIM, font='Arial')
            if rep:
                para.add_run('\n'); s.add_run(para, ('%s : ' % qui) if qui else '', bold=True, size=8.5, color=s.C_DIM, font='Arial Narrow')
                s.add_run(para, rep, italic=True, size=9, color=s.C_GREY)
        elif k == 'rattrapage':
            s.recovery_block(b[1], b[2], b[3], b[4], b[5])
        elif k == 'quiz':
            s.quiz_block(b[1], b[2], b[3], b[4], b[5])
        elif k == 'dm':
            s.dm_block(b[1], b[2], b[3], b[4])
        elif k == 'photo':
            head = doc.add_paragraph(); s.set_spacing(head, before=60, after=10)
            s.set_indent(head, left=200); s.left_border(head, '3a4a6a', sz=14); s.shading(head, 'F0F4FA')
            s.add_run(head, 'Photo %d - %s' % (b[1] + 1, b[2]), bold=True, size=9.5, color=s.C_BLUE, font='Arial')
            s.add_run(head, '   %d j’aime' % b[3], size=8, color=s.C_DIM, font='Arial')
            for ligne in b[4]:
                mp = doc.add_paragraph(); s.set_spacing(mp, before=0, after=8, line=260)
                s.set_indent(mp, left=400); s.shading(mp, 'F7F9FD')
                s.add_run(mp, ligne, size=9.5)
        elif k == 'carte':
            para = doc.add_paragraph(); s.set_spacing(para, before=40, after=40)
            s.set_indent(para, left=200); s.shading(para, 'F7F9FD'); s.left_border(para, '3a4a6a', sz=10)
            s.add_run(para, b[1], bold=True, size=9.5, color=s.C_BLUE, font='Arial')
            para.add_run('\n'); s.add_run(para, b[2], size=9.5)
        elif k == 'compte':
            para = doc.add_paragraph(); s.set_spacing(para, before=0, after=20)
            s.set_indent(para, left=300)
            s.add_run(para, b[1] + '  →  ', size=10, color=s.C_DIM, font='Consolas')
            s.add_run(para, b[2], bold=True, size=10)
        elif k == 'fin':
            s.ending_block(b[1], b[2], *b[3])
        elif k == 'ressource':
            para = doc.add_paragraph(); s.set_spacing(para, before=0, after=20)
            s.set_indent(para, left=300)
            s.add_run(para, b[1] + '   ', bold=True, size=10)
            s.add_run(para, b[2], size=10, color=s.C_GREY)
    s.footnote_line('Retrouver Clara - document produit automatiquement à partir des fichiers du jeu.')
    doc.save('doc_textes_partenaires.docx')

# ═══════════════════════════════════════════════════════════════════════════
#  RENDU HTML
# ═══════════════════════════════════════════════════════════════════════════
CSS = """
*{box-sizing:border-box}
body{margin:0;background:#eeedea;color:#1a1a1e;font-family:'Jost',system-ui,-apple-system,'Segoe UI',sans-serif;font-size:15px;line-height:1.62}
.page{max-width:820px;margin:0 auto;padding:40px 26px 70px;background:#fff;box-shadow:0 0 40px rgba(0,0,0,.06)}
.cover{text-align:center;padding:40px 0 34px;border-bottom:2px solid #1a1a1e;margin-bottom:8px}
.cover .ey{font-size:11px;letter-spacing:.2em;text-transform:uppercase;color:#7a7a83}
.cover h1{font-family:Georgia,serif;font-style:italic;font-weight:400;font-size:40px;margin:14px 0 12px}
.cover .su{color:#55555e;font-size:15px;white-space:pre-line}
.cover .me{font-size:11px;letter-spacing:.14em;text-transform:uppercase;color:#9a9aa2;margin-top:14px}
h2.part{margin:52px 0 4px;padding:16px 18px;background:#1a1a1e;color:#fff;border-radius:3px;font-size:19px;font-weight:500}
h2.part span{display:block;font-size:11px;letter-spacing:.18em;text-transform:uppercase;color:#a9a9b2;margin-bottom:5px}
h3{font-family:Georgia,serif;font-weight:400;font-size:21px;margin:30px 0 8px}
h4{font-size:15px;margin:22px 0 6px;color:#3a3a44}
p{margin:0 0 10px;max-width:70ch}
.note{font-size:13.5px;color:#6a6a73;font-style:italic;border-left:2px solid #d8d7d2;padding-left:12px;margin:12px 0}
ul{margin:6px 0 12px;padding-left:20px} li{margin-bottom:5px}
table{border-collapse:collapse;margin:12px 0 18px;font-size:14px;width:100%}
th,td{border:1px solid #ddddd8;padding:7px 10px;text-align:left}
th{background:#f4f3f0;font-weight:600;font-size:13px}
.dial{margin:0 0 6px;display:flex;gap:10px;align-items:baseline}
.dial .who{flex-shrink:0;min-width:74px;font-size:11.5px;font-weight:700;text-transform:uppercase;letter-spacing:.04em}
.dial .txt{white-space:pre-line}
.dial.narr .txt{color:#8a8a92;font-style:italic}
.prompt{font-size:13px;color:#8a8a92;font-style:italic;margin:14px 0 6px}
.pensee{background:#f2f8f2;border-left:3px solid #4dcc70;padding:9px 13px;margin:8px 0;font-style:italic;color:#2a7a50;font-size:14px}
.choix{border-left:3px solid;padding:9px 13px;margin:0 0 7px;border-radius:0 3px 3px 0}
.choix .k{font-weight:700;font-size:12px;margin-right:7px}
.choix .n{display:block;font-size:12.5px;color:#7a7a83;font-style:italic;margin-top:3px}
.choix .r{display:block;font-size:13px;color:#55555e;margin-top:4px}
.choix.good{border-color:#2a7a50;background:#f4faf7}
.choix.ok{border-color:#c09030;background:#fdfaf0}
.choix.bad{border-color:#c04030;background:#fdf4f4}
.choix.neutral{border-color:#888;background:#fafafa}
.recov{background:#fffbf0;border-left:3px solid #e0c080;padding:10px 13px;margin:6px 0 12px;font-size:14px}
.recov .lab{font-size:12px;color:#a08030;font-style:italic;display:block;margin-bottom:6px}
.quiz{border:1px solid #dcdcd6;border-radius:4px;padding:13px 15px;margin:12px 0}
.quiz .ou{font-size:11.5px;letter-spacing:.1em;text-transform:uppercase;color:#8a8a92;margin-bottom:6px}
.quiz .q{font-weight:600;margin-bottom:8px}
.quiz ol{margin:0 0 8px;padding-left:20px}
.quiz li.ok{color:#2a7a50;font-weight:600}
.quiz .exp{font-size:13px;color:#6a6a73;font-style:italic;border-top:1px solid #eeeee8;padding-top:8px}
.dm{border-left:3px solid #3a4a6a;background:#f7f9fd;padding:10px 13px;margin:0 0 10px}
.dm .hd{font-size:13px;font-weight:700;color:#2a4a7a}
.dm .hd .t{font-weight:400;color:#9a9aa2;margin-left:8px;font-size:12px}
.dm .hd .ty{font-style:italic;color:#c09030;font-size:12px;margin-left:8px}
.dm .m{font-size:13.5px;margin-top:5px;color:#3a3a44}
.dm.amie{border-color:#44aa66;background:#f4faf6}
.carte{border-left:3px solid #3a4a6a;background:#f7f9fd;padding:9px 13px;margin:0 0 8px}
.carte b{color:#2a4a7a;font-size:14px}
.carte span{display:block;font-size:13.5px;color:#4a4a52;margin-top:3px}
.compte{font-family:ui-monospace,Menlo,monospace;font-size:13.5px;margin:0 0 5px}
.compte b{font-family:'Jost',sans-serif;font-size:14.5px}
.fin{border-left:3px solid;padding:11px 14px;margin:0 0 11px;border-radius:0 3px 3px 0}
.fin .t{font-weight:700;font-size:14px;margin-bottom:5px}
.fin p{font-size:13.5px;margin:0 0 5px;white-space:pre-line}
.sep{text-align:center;font-size:12px;letter-spacing:.14em;text-transform:uppercase;color:#9a9aa2;margin:20px 0 12px}
.res{display:flex;gap:14px;font-size:14px;margin-bottom:5px}
.res b{min-width:120px}
footer{margin-top:40px;padding-top:14px;border-top:1px solid #e6e5e0;font-size:12px;color:#9a9aa2;text-align:center}
@media print{body{background:#fff}.page{box-shadow:none;max-width:none;padding:0}
  h2.part{background:#1a1a1e !important;-webkit-print-color-adjust:exact;print-color-adjust:exact}
  .dm,.carte,.quiz,.fin,.choix{page-break-inside:avoid}}
"""
SPKCOL = {'leo': '#1a3a6a', 'ines': '#4a1a6a', 'clara': '#6a1a2a', 'lea': '#448844',
          'enzo': '#1a5a28', 'jade': '#7a3a00', 'théo': '#2a4a7a', 'marine': '#7a1a4a',
          'kevin': '#555', 'aunt': '#5a3510', 'thomas': '#1a3a6a', 'narr': '#9a9aa2', 'anon': '#8a8a92'}

def e(x): return _h.escape(str(x)) if x is not None else ''

def rendu_html():
    o = ['<!DOCTYPE html>', '<html lang="fr"><head><meta charset="utf-8">',
         '<meta name="viewport" content="width=device-width, initial-scale=1.0">',
         '<title>Retrouver Clara - textes complets</title>',
         '<link href="https://fonts.googleapis.com/css2?family=Jost:wght@300;400;500;600;700&display=swap" rel="stylesheet">',
         '<style>%s</style></head><body><div class="page">' % CSS]
    for b in B:
        k = b[0]
        if k == 'cover':
            o.append('<div class="cover"><div class="ey">Jeu pédagogique sur le harcèlement scolaire</div>'
                     '<h1>%s</h1><div class="su">%s</div><div class="me">%s</div></div>'
                     % (e(b[1]), e(b[2]), e(b[3])))
        elif k == 'h1':
            o.append('<h2 class="part"><span>%s</span>%s</h2>' % (e(b[1]), e(b[2])))
        elif k == 'h2':  o.append('<h3>%s</h3>' % e(b[1]))
        elif k == 'h3':  o.append('<h4>%s</h4>' % e(b[1]))
        elif k == 'p':   o.append('<p>%s</p>' % e(b[1]).replace('\n', '<br>'))
        elif k == 'note': o.append('<div class="note">%s</div>' % e(b[1]).replace('\n', '<br>'))
        elif k == 'bullet': o.append('<ul><li>%s</li></ul>' % e(b[1]))
        elif k == 'sep': o.append('<div class="sep">- %s -</div>' % e(b[1]))
        elif k == 'prompt': o.append('<div class="prompt">%s</div>' % e(b[1]))
        elif k == 'pensee': o.append('<div class="pensee">%s</div>' % e(b[1]))
        elif k == 'table':
            o.append('<table><tr>%s</tr>' % ''.join('<th>%s</th>' % e(x) for x in b[1]))
            for ln in b[2]:
                o.append('<tr>%s</tr>' % ''.join('<td>%s</td>' % e(x) for x in ln))
            o.append('</table>')
        elif k == 'dial':
            col = SPKCOL.get(b[1], '#55555e')
            narr = ' narr' if (len(b) > 4 and b[4]) else ''
            o.append('<div class="dial%s"><span class="who" style="color:%s">%s</span>'
                     '<span class="txt">%s</span></div>' % (narr, col, e(b[2]), e(b[3])))
        elif k == 'choix':
            kind, lettre, txt, note, rep, qui = b[1], b[2], b[3], b[4], b[5], b[6]
            h = '<div class="choix %s"><span class="k">%s</span>%s' % (kind, e(lettre), e(txt))
            if note: h += '<span class="n">%s</span>' % e(note)
            if rep:  h += '<span class="r">%s%s</span>' % (('<b>%s :</b> ' % e(qui)) if qui else '', e(rep))
            o.append(h + '</div>')
        elif k == 'rattrapage':
            o.append('<div class="recov"><span class="lab">Rattrapage possible (si mauvaise réponse)</span>'
                     '<p><b>La tante :</b> %s</p><p>✓ %s<br><i>%s</i></p><p>✗ %s<br><i>%s</i></p></div>'
                     % (e(b[1]), e(b[2]), e(b[3]), e(b[4]), e(b[5])))
        elif k == 'quiz':
            opts = ''.join('<li class="%s">%s</li>' % ('ok' if i == b[4] else '', e(x))
                           for i, x in enumerate(b[3]))
            o.append('<div class="quiz"><div class="ou">%s</div><div class="q">%s</div>'
                     '<ol>%s</ol><div class="exp">%s</div></div>' % (e(b[1]), e(b[2]), opts, e(b[5])))
        elif k == 'dm':
            amie = ' amie' if (len(b) > 5 and b[5]) else ''
            ty = '<span class="ty">%s</span>' % e(b[3]) if b[3] else ''
            msgs = ''.join('<div class="m">« %s »</div>' % e(m) for m in b[4])
            o.append('<div class="dm%s"><div class="hd">%s<span class="t">%s</span>%s</div>%s</div>'
                     % (amie, e(b[1]), e(b[2]), ty, msgs))
        elif k == 'photo':
            msgs = ''.join('<div class="m">%s</div>' % e(m) for m in b[4])
            o.append('<div class="dm"><div class="hd">Photo %d - %s<span class="t">%d j’aime</span></div>%s</div>'
                     % (b[1] + 1, e(b[2]), b[3], msgs))
        elif k == 'carte':
            o.append('<div class="carte"><b>%s</b><span>%s</span></div>' % (e(b[1]), e(b[2])))
        elif k == 'compte':
            o.append('<div class="compte">%s → <b>%s</b></div>' % (e(b[1]), e(b[2])))
        elif k == 'fin':
            col = {'good': '#2a7a50', 'ok': '#c09030', 'bad': '#c04030', 'neutral': '#888'}[b[1]]
            body = ''.join('<p>%s</p>' % e(x) for x in b[3])
            o.append('<div class="fin" style="border-color:%s;background:%s"><div class="t" style="color:%s">%s</div>%s</div>'
                     % (col, {'good': '#f4faf7', 'ok': '#fdfaf0', 'bad': '#fdf4f4', 'neutral': '#fafafa'}[b[1]],
                        col, e(b[2]), body))
        elif k == 'ressource':
            o.append('<div class="res"><b>%s</b><span>%s</span></div>' % (e(b[1]), e(b[2])))
    o.append('<footer>Retrouver Clara - document produit automatiquement à partir des fichiers du jeu '
             '(<code>node dump_textes.js</code> puis <code>python gen_doc.py</code>).</footer>')
    o.append('</div></body></html>')
    open('doc_textes_partenaires.html', 'w', encoding='utf-8', newline='\n').write('\n'.join(o))

rendu_docx()
rendu_html()
print('écrits : doc_textes_partenaires.docx et doc_textes_partenaires.html (%d blocs)' % len(B))
